"""考试报表：按权限范围聚合考试成绩，支持三视角查询与多模板 Excel 导出。"""

from __future__ import annotations

from collections import defaultdict
from io import BytesIO
from typing import Any, Optional

from django.db.models import Sum
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from apps.authorization.engine import authorize, scope_filter
from apps.submissions.models import Submission
from apps.tasks.models import Task, TaskAssignment, TaskQuiz
from apps.users.models import User
from core.base_service import BaseService

SUBMISSION_STATUS_LABELS = {
    'IN_PROGRESS': '答题中',
    'SUBMITTED': '已提交',
    'GRADING': '待评分',
    'GRADED': '已评分',
}

PASS_PASSED = '及格'
PASS_FAILED = '不及格'
PASS_PENDING = '待判定'
PASS_ABSENT = '未参加'

# 无排名时排序用，保证未出分排在最后
_RANK_SORT_MISSING = 10**9


class ExamReportService(BaseService):
    """以任务内考试节点（TaskQuiz + EXAM）为单位聚合学员成绩。"""

    def get_report(self, filters: dict[str, Any]) -> dict[str, Any]:
        view = filters.get('view') or 'detail'
        selected_exam_id = filters.get('exam_id')
        page = int(filters.get('page') or 1)
        page_size = int(filters.get('page_size') or 10)

        ctx = self._load_context(filters)
        if view == 'student':
            rows, total = self._page_student_rows(ctx, page, page_size)
        elif view == 'exam':
            rows, total = self._page_detail_rows(
                ctx,
                page,
                page_size,
                exam_ids=[selected_exam_id] if selected_exam_id else [],
                sort_exam_view=True,
            )
        else:
            rows, total = self._page_detail_rows(ctx, page, page_size)

        total_pages = max(1, (total + page_size - 1) // page_size) if page_size else 1
        if page > total_pages:
            page = total_pages

        return {
            'view': view,
            'selected_exam_id': selected_exam_id,
            'summary': {
                'student_count': ctx['summary_student_count'],
                'exam_count': len(ctx['exams']),
                'record_count': ctx['summary_record_count'],
            },
            'filters': {
                'exams': [
                    {
                        'id': exam['id'],
                        'label': exam['label'],
                        'exam_title': exam['exam_title'],
                        'task_title': exam['task_title'],
                    }
                    for exam in ctx['all_exams']
                ],
                'students': [
                    {
                        'id': student['id'],
                        'name': student['name'],
                        'employee_id': student['employee_id'],
                    }
                    for student in ctx['all_students']
                ],
                'departments': ctx['departments'],
            },
            'exams': [
                {
                    'id': exam['id'],
                    'label': exam['label'],
                    'exam_title': exam['exam_title'],
                    'task_title': exam['task_title'],
                    'total_score': exam['total_score'],
                    'pass_score': exam['pass_score'],
                }
                for exam in ctx['exams']
            ],
            'rows': rows,
            'pagination': {
                'count': total,
                'page': page,
                'page_size': page_size,
                'total_pages': total_pages,
            },
        }

    def export_report(self, filters: dict[str, Any], template: str) -> bytes:
        ctx = self._load_context(filters)
        workbook = Workbook()
        workbook.remove(workbook.active)

        if template == 'detail':
            records = self._materialize_records(ctx, ctx['filtered_student_ids'])
            self._write_detail_sheet(workbook, records)
        elif template == 'student_summary':
            data = {
                'exams': ctx['exams'],
                'students': [
                    ctx['student_map'][sid]
                    for sid in ctx['eligible_student_ids']
                ],
                'records': self._materialize_records(ctx, ctx['filtered_student_ids']),
            }
            self._write_student_sheet(workbook, data)
        elif template == 'exam_summary':
            records = self._materialize_records(ctx, ctx['filtered_student_ids'])
            self._write_exam_sheet(workbook, {'records': records})
        else:
            raise ValueError(f'不支持的导出模板: {template}')

        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    # ------------------------------------------------------------------
    # 上下文：权限学员、考试、分配驱动的参与对、轻量排名
    # ------------------------------------------------------------------

    def _load_context(self, filters: dict[str, Any]) -> dict[str, Any]:
        student_qs = self._accessible_students().select_related('department', 'mentor')
        all_students = [
            {
                'id': user.id,
                'name': user.username,
                'employee_id': user.employee_id or '',
                'department_id': user.department_id,
                'department_name': user.department.name if user.department_id else '',
                'mentor_name': user.mentor.username if user.mentor_id else '',
                'avatar_key': user.avatar_key or '',
            }
            for user in student_qs.order_by('employee_id', 'id')
        ]
        student_map = {item['id']: item for item in all_students}
        student_ids = list(student_map.keys())

        departments = sorted(
            {
                item['department_id']: {
                    'id': item['department_id'],
                    'name': item['department_name'],
                }
                for item in all_students
                if item['department_id']
            }.values(),
            key=lambda item: item['name'],
        )

        all_exams = self._load_exams(student_ids)
        exam_map = {exam['id']: exam for exam in all_exams}
        filtered_exam_ids = self._apply_exam_filters(list(exam_map.keys()), filters)
        exams = [exam_map[eid] for eid in filtered_exam_ids if eid in exam_map]
        filtered_student_ids = self._apply_student_filters(student_ids, student_map, filters)
        filtered_student_id_set = set(filtered_student_ids)

        # 以任务分配驱动：只处理真实 (考试任务, 学员) 对，避免 E×S 空扫
        assignees_by_task: dict[int, list[int]] = defaultdict(list)
        if student_ids and exams:
            task_ids = {exam['task_id'] for exam in exams}
            for row in TaskAssignment.objects.filter(
                task_id__in=task_ids,
                assignee_id__in=student_ids,
            ).only('task_id', 'assignee_id'):
                assignees_by_task[row.task_id].append(row.assignee_id)

        cohort_student_ids = list({
            sid
            for exam in exams
            for sid in assignees_by_task.get(exam['task_id'], [])
        })
        by_pair = self._load_submission_pairs(
            student_ids=cohort_student_ids,
            exam_ids=[exam['id'] for exam in exams],
        )

        # 每场考试在权限全量参与者上轻量算分/排名；筛选只影响返回行
        rank_by_exam: dict[int, dict[int, Optional[int]]] = {}
        score_info_by_exam: dict[int, dict[int, dict[str, Any]]] = {}
        for exam in exams:
            assignee_ids = assignees_by_task.get(exam['task_id'], [])
            info_by_student: dict[int, dict[str, Any]] = {}
            rank_inputs: list[dict[str, Any]] = []
            for sid in assignee_ids:
                student = student_map.get(sid)
                if not student:
                    continue
                submission = by_pair.get((exam['id'], sid))
                score, pass_status, submission_status, attempt_number, time_spent = (
                    self._extract_submission_fields(submission, exam['pass_score'])
                )
                info_by_student[sid] = {
                    'score': score,
                    'pass_status': pass_status,
                    'submission_status': submission_status,
                    'attempt_number': attempt_number,
                    'time_spent_minutes': time_spent,
                    'employee_id': student['employee_id'],
                }
                if score is not None:
                    rank_inputs.append({
                        'student_id': sid,
                        'score': score,
                        'employee_id': student['employee_id'],
                    })
            rank_by_exam[exam['id']] = self._compute_rank_map(rank_inputs)
            score_info_by_exam[exam['id']] = info_by_student

        # 筛选后仍有分配记录的学员 / 明细条数（不物化完整 record）
        eligible_student_ids: list[int] = []
        summary_record_count = 0
        seen_students: set[int] = set()
        for exam in exams:
            for sid in assignees_by_task.get(exam['task_id'], []):
                if sid not in filtered_student_id_set:
                    continue
                summary_record_count += 1
                if sid not in seen_students:
                    seen_students.add(sid)
                    eligible_student_ids.append(sid)
        eligible_student_ids.sort(key=lambda sid: student_map[sid]['employee_id'])

        return {
            'all_students': all_students,
            'all_exams': all_exams,
            'student_map': student_map,
            'exams': exams,
            'departments': departments,
            'filtered_student_ids': filtered_student_ids,
            'filtered_student_id_set': filtered_student_id_set,
            'eligible_student_ids': eligible_student_ids,
            'assignees_by_task': assignees_by_task,
            'rank_by_exam': rank_by_exam,
            'score_info_by_exam': score_info_by_exam,
            'summary_student_count': len(eligible_student_ids),
            'summary_record_count': summary_record_count,
        }

    def _accessible_students(self):
        """按 task.analytics.view 取学员范围。

        有分析权限：走 scope_filter（导师/室经理各自范围，管理员 ALL）。
        无分析权限：仅 dashboard.admin.view 可回退全员；导师/室经理返回空集，避免越权。
        """
        base = User.objects.filter(
            is_active=True,
            roles__code='STUDENT',
        ).exclude(is_superuser=True).distinct()

        if authorize('task.analytics.view', self.request).allowed:
            return scope_filter(
                'task.analytics.view',
                self.request,
                base_queryset=base,
                resource_model=User,
            )
        if authorize('dashboard.admin.view', self.request).allowed:
            return base
        return base.none()

    def _load_exams(self, student_ids: list[int]) -> list[dict[str, Any]]:
        if not student_ids:
            return []

        # 学员被分配的任务 ∩ 当前用户 task.view 可见任务，避免泄露他人创建的不可见任务
        assigned_task_ids = (
            TaskAssignment.objects.filter(assignee_id__in=student_ids)
            .values_list('task_id', flat=True)
            .distinct()
        )
        visible_task_ids = scope_filter(
            'task.view',
            self.request,
            resource_model=Task,
            base_queryset=Task.objects.filter(id__in=assigned_task_ids),
        ).values_list('id', flat=True)
        quizzes = (
            TaskQuiz.objects.filter(
                task_id__in=visible_task_ids,
                quiz__quiz_type='EXAM',
            )
            .select_related('task', 'quiz')
            .annotate(total_score_sum=Sum('quiz__quiz_questions__score'))
            .order_by('-task__deadline', 'task_id', 'order', 'id')
        )

        exams: list[dict[str, Any]] = []
        for item in quizzes:
            total_score = float(item.total_score_sum or 0)
            pass_score = float(item.quiz.pass_score) if item.quiz.pass_score is not None else None
            exams.append(
                {
                    'id': item.id,
                    'task_id': item.task_id,
                    'task_title': item.task.title,
                    'exam_title': item.quiz.title,
                    'label': f'{item.quiz.title} · {item.task.title}',
                    'deadline': item.task.deadline.isoformat() if item.task.deadline else None,
                    'pass_score': pass_score,
                    'total_score': total_score,
                }
            )
        return exams

    def _apply_student_filters(
        self,
        student_ids: list[int],
        student_map: dict[int, dict[str, Any]],
        filters: dict[str, Any],
    ) -> list[int]:
        result = student_ids
        student_id = filters.get('student_id')
        if student_id:
            result = [sid for sid in result if sid == student_id]

        department_id = filters.get('department_id')
        if department_id:
            result = [
                sid for sid in result
                if student_map[sid]['department_id'] == department_id
            ]

        search = (filters.get('search') or '').strip().lower()
        if search:
            result = [
                sid for sid in result
                if search in student_map[sid]['name'].lower()
                or search in student_map[sid]['employee_id'].lower()
            ]
        return result

    def _apply_exam_filters(self, exam_ids: list[int], filters: dict[str, Any]) -> list[int]:
        exam_id = filters.get('exam_id')
        if not exam_id:
            return exam_ids
        return [eid for eid in exam_ids if eid == exam_id]

    def _load_submission_pairs(
        self,
        *,
        student_ids: list[int],
        exam_ids: list[int],
    ) -> dict[tuple[int, int], Submission]:
        if not student_ids or not exam_ids:
            return {}

        submissions = (
            Submission.objects.filter(
                task_quiz_id__in=exam_ids,
                user_id__in=student_ids,
            )
            .only(
                'id',
                'task_quiz_id',
                'user_id',
                'status',
                'obtained_score',
                'attempt_number',
                'started_at',
                'submitted_at',
                'created_at',
            )
            .order_by('task_quiz_id', 'user_id', '-attempt_number', '-created_at', '-id')
        )
        by_pair: dict[tuple[int, int], Submission] = {}
        for submission in submissions:
            key = (submission.task_quiz_id, submission.user_id)
            if key not in by_pair:
                by_pair[key] = submission
        return by_pair

    def _extract_submission_fields(
        self,
        submission: Optional[Submission],
        pass_score: Optional[float],
    ) -> tuple[Optional[float], str, str, Optional[int], Optional[int]]:
        if submission is None:
            return None, PASS_ABSENT, '未参加', None, None

        submission_status = SUBMISSION_STATUS_LABELS.get(submission.status, submission.status)
        attempt_number = submission.attempt_number
        # SUBMITTED=纯客观自动终分；GRADED=阅卷终分；GRADING 仅客观题小计，不当终分
        if (
            submission.status in {'SUBMITTED', 'GRADED'}
            and submission.obtained_score is not None
        ):
            score = float(submission.obtained_score)
        else:
            score = None
        pass_status = self._resolve_pass_status(score, pass_score, submission.status)
        if submission.submitted_at and submission.started_at:
            seconds = max(0.0, (submission.submitted_at - submission.started_at).total_seconds())
            time_spent_minutes = int(seconds / 60)
        else:
            time_spent_minutes = None
        return score, pass_status, submission_status, attempt_number, time_spent_minutes

    @staticmethod
    def _resolve_pass_status(
        score: Optional[float],
        pass_score: Optional[float],
        submission_status: str,
    ) -> str:
        if score is None:
            return PASS_PENDING
        if pass_score is None:
            return PASS_PENDING
        return PASS_PASSED if score >= pass_score else PASS_FAILED

    @staticmethod
    def _compute_rank_map(scored_rows: list[dict[str, Any]]) -> dict[int, Optional[int]]:
        scored_rows = sorted(
            scored_rows,
            key=lambda item: (-item['score'], item['employee_id'], item['student_id']),
        )
        rank_map: dict[int, Optional[int]] = {}
        last_score = None
        last_rank = 0
        for index, item in enumerate(scored_rows, start=1):
            if last_score is None or item['score'] != last_score:
                last_rank = index
                last_score = item['score']
            rank_map[item['student_id']] = last_rank
        return rank_map

    def _compose_record(
        self,
        exam: dict[str, Any],
        student: dict[str, Any],
        info: dict[str, Any],
        rank: Optional[int],
    ) -> dict[str, Any]:
        return {
            'student_id': student['id'],
            'student_name': student['name'],
            'employee_id': student['employee_id'],
            'department_name': student['department_name'],
            'mentor_name': student['mentor_name'],
            'avatar_key': student['avatar_key'],
            'exam_id': exam['id'],
            'exam_title': exam['exam_title'],
            'task_id': exam['task_id'],
            'task_title': exam['task_title'],
            'score': info['score'],
            'total_score': exam['total_score'],
            'pass_score': exam['pass_score'],
            'rank': rank,
            'pass_status': info['pass_status'],
            'submission_status': info['submission_status'],
            'attempt_number': info['attempt_number'],
            'time_spent_minutes': info['time_spent_minutes'],
        }

    def _iter_index(
        self,
        ctx: dict[str, Any],
        *,
        student_id_set: set[int],
        exam_ids: Optional[list[int]] = None,
    ):
        """产出 (exam, student) 轻量索引，用于排序与分页，不物化完整行。"""
        exams = ctx['exams']
        if exam_ids is not None:
            exam_id_set = set(exam_ids)
            exams = [exam for exam in exams if exam['id'] in exam_id_set]

        for exam in exams:
            info_by_student = ctx['score_info_by_exam'].get(exam['id'], {})
            ranks = ctx['rank_by_exam'].get(exam['id'], {})
            for sid in ctx['assignees_by_task'].get(exam['task_id'], []):
                if sid not in student_id_set:
                    continue
                student = ctx['student_map'].get(sid)
                if not student or sid not in info_by_student:
                    continue
                rank = ranks.get(sid)
                yield {
                    'exam': exam,
                    'student': student,
                    'info': info_by_student[sid],
                    'rank': rank,
                    'sort_rank': rank if rank is not None else _RANK_SORT_MISSING,
                    'employee_id': student['employee_id'],
                    'exam_title': exam['exam_title'],
                    'task_title': exam['task_title'],
                }

    def _materialize_records(
        self,
        ctx: dict[str, Any],
        student_ids: list[int],
        exam_ids: Optional[list[int]] = None,
    ) -> list[dict[str, Any]]:
        index = list(
            self._iter_index(
                ctx,
                student_id_set=set(student_ids),
                exam_ids=exam_ids,
            )
        )
        index.sort(
            key=lambda item: (
                item['exam_title'],
                item['task_title'],
                item['sort_rank'],
                item['employee_id'],
            )
        )
        return [
            self._compose_record(item['exam'], item['student'], item['info'], item['rank'])
            for item in index
        ]

    def _page_detail_rows(
        self,
        ctx: dict[str, Any],
        page: int,
        page_size: int,
        *,
        exam_ids: Optional[list[int]] = None,
        sort_exam_view: bool = False,
    ) -> tuple[list[dict[str, Any]], int]:
        if exam_ids is not None and not exam_ids:
            return [], 0

        index = list(
            self._iter_index(
                ctx,
                student_id_set=ctx['filtered_student_id_set'],
                exam_ids=exam_ids,
            )
        )
        if sort_exam_view:
            index.sort(key=lambda item: (item['sort_rank'], item['employee_id']))
        else:
            index.sort(
                key=lambda item: (
                    item['exam_title'],
                    item['task_title'],
                    item['sort_rank'],
                    item['employee_id'],
                )
            )

        total = len(index)
        if total == 0:
            return [], 0

        total_pages = max(1, (total + page_size - 1) // page_size) if page_size else 1
        if page > total_pages:
            page = total_pages
        start = (page - 1) * page_size
        page_items = index[start:start + page_size]
        rows = [
            self._compose_record(item['exam'], item['student'], item['info'], item['rank'])
            for item in page_items
        ]
        return rows, total

    def _page_student_rows(
        self,
        ctx: dict[str, Any],
        page: int,
        page_size: int,
    ) -> tuple[list[dict[str, Any]], int]:
        eligible = ctx['eligible_student_ids']
        total = len(eligible)
        if total == 0:
            return [], 0

        total_pages = max(1, (total + page_size - 1) // page_size) if page_size else 1
        if page > total_pages:
            page = total_pages
        start = (page - 1) * page_size
        page_student_ids = eligible[start:start + page_size]

        exam_ids = [exam['id'] for exam in ctx['exams']]
        rows: list[dict[str, Any]] = []
        for sid in page_student_ids:
            student = ctx['student_map'][sid]
            exam_scores: dict[str, Optional[float]] = {str(eid): None for eid in exam_ids}
            exam_pass: dict[str, str] = {str(eid): PASS_ABSENT for eid in exam_ids}
            scores: list[float] = []
            passed_count = 0
            for exam in ctx['exams']:
                info = ctx['score_info_by_exam'].get(exam['id'], {}).get(sid)
                if info is None:
                    continue
                key = str(exam['id'])
                exam_scores[key] = info['score']
                exam_pass[key] = info['pass_status']
                if info['score'] is not None:
                    scores.append(info['score'])
                if info['pass_status'] == PASS_PASSED:
                    passed_count += 1
            scored_count = len(scores)
            rows.append(
                {
                    'student_id': student['id'],
                    'student_name': student['name'],
                    'employee_id': student['employee_id'],
                    'department_name': student['department_name'],
                    'mentor_name': student['mentor_name'],
                    'avatar_key': student['avatar_key'],
                    'scored_count': scored_count,
                    'average_score': round(sum(scores) / scored_count, 1) if scored_count else None,
                    'passed_count': passed_count,
                    'pass_ratio': f'{passed_count}/{scored_count}' if scored_count else '0/0',
                    'exam_scores': exam_scores,
                    'exam_pass': exam_pass,
                }
            )
        return rows, total

    def _build_student_rows(self, data: dict[str, Any]) -> list[dict[str, Any]]:
        """导出学员汇总用：基于已物化 records。"""
        exam_ids = [exam['id'] for exam in data['exams']]
        grouped: dict[int, list[dict[str, Any]]] = defaultdict(list)
        for record in data['records']:
            grouped[record['student_id']].append(record)

        rows: list[dict[str, Any]] = []
        for student in data['students']:
            student_records = grouped.get(student['id'], [])
            exam_scores: dict[str, Optional[float]] = {
                str(exam_id): None for exam_id in exam_ids
            }
            exam_pass: dict[str, str] = {
                str(exam_id): PASS_ABSENT for exam_id in exam_ids
            }
            scores: list[float] = []
            passed_count = 0
            for record in student_records:
                key = str(record['exam_id'])
                exam_scores[key] = record['score']
                exam_pass[key] = record['pass_status']
                if record['score'] is not None:
                    scores.append(record['score'])
                if record['pass_status'] == PASS_PASSED:
                    passed_count += 1

            scored_count = len(scores)
            rows.append(
                {
                    'student_id': student['id'],
                    'student_name': student['name'],
                    'employee_id': student['employee_id'],
                    'department_name': student['department_name'],
                    'mentor_name': student['mentor_name'],
                    'avatar_key': student['avatar_key'],
                    'scored_count': scored_count,
                    'average_score': round(sum(scores) / scored_count, 1) if scored_count else None,
                    'passed_count': passed_count,
                    'pass_ratio': f'{passed_count}/{scored_count}' if scored_count else '0/0',
                    'exam_scores': exam_scores,
                    'exam_pass': exam_pass,
                }
            )

        rows.sort(key=lambda item: item['employee_id'])
        return rows

    # ------------------------------------------------------------------
    # Excel
    # ------------------------------------------------------------------

    @staticmethod
    def _excel_safe(value: Any) -> Any:
        """避免用户可控文本以 =+-@ 等开头被 Excel 当公式执行。"""
        if isinstance(value, str) and value and value[0] in ('=', '+', '-', '@', '\t', '\r'):
            return f"'{value}"
        return value

    def _write_cells(self, sheet, row_index: int, values: list[Any]) -> None:
        for col, value in enumerate(values, start=1):
            sheet.cell(row=row_index, column=col, value=self._excel_safe(value))

    @staticmethod
    def _style_header(sheet, headers: list[str]) -> None:
        fill = PatternFill('solid', fgColor='183B5B')
        font = Font(color='FFFFFF', bold=True)
        alignment = Alignment(horizontal='center', vertical='center')
        for col, header in enumerate(headers, start=1):
            safe = ExamReportService._excel_safe(header)
            cell = sheet.cell(row=1, column=col, value=safe)
            cell.fill = fill
            cell.font = font
            cell.alignment = alignment
        sheet.freeze_panes = 'A2'

    def _write_detail_sheet(self, workbook: Workbook, records: list[dict[str, Any]]) -> None:
        sheet = workbook.create_sheet('考试明细')
        headers = [
            '学员姓名', '工号', '部门', '导师', '考试', '所属任务',
            '成绩', '满分', '排名', '是否及格', '作答状态', '考试用时(分钟)',
        ]
        self._style_header(sheet, headers)
        for row_index, item in enumerate(records, start=2):
            self._write_cells(sheet, row_index, [
                item['student_name'],
                item['employee_id'],
                item['department_name'],
                item['mentor_name'],
                item['exam_title'],
                item['task_title'],
                self._format_score(item['score']),
                self._format_score(item['total_score']),
                item['rank'] if item['rank'] is not None else '',
                item['pass_status'],
                item['submission_status'],
                item['time_spent_minutes'] if item['time_spent_minutes'] is not None else '',
            ])

    def _write_student_sheet(self, workbook: Workbook, data: dict[str, Any]) -> None:
        sheet = workbook.create_sheet('学员汇总')
        exams = data['exams']
        headers = ['学员姓名', '工号', '部门', '导师', '已考场次', '平均分', '及格场次']
        headers.extend(exam['label'] for exam in exams)
        self._style_header(sheet, headers)

        rows = self._build_student_rows(data)
        for row_index, item in enumerate(rows, start=2):
            values = [
                item['student_name'],
                item['employee_id'],
                item['department_name'],
                item['mentor_name'],
                item['scored_count'],
                self._format_score(item['average_score']),
                item['pass_ratio'],
            ]
            for exam in exams:
                values.append(self._format_score(item['exam_scores'].get(str(exam['id']))))
            self._write_cells(sheet, row_index, values)

    def _write_exam_sheet(self, workbook: Workbook, data: dict[str, Any]) -> None:
        sheet = workbook.create_sheet('考试汇总')
        headers = [
            '考试', '所属任务', '排名', '学员姓名', '工号', '部门',
            '成绩', '满分', '是否及格', '作答状态', '考试用时(分钟)',
        ]
        self._style_header(sheet, headers)

        records = sorted(
            data['records'],
            key=lambda item: (
                item['exam_title'],
                item['task_title'],
                item['rank'] if item['rank'] is not None else _RANK_SORT_MISSING,
                item['employee_id'],
            ),
        )
        for row_index, item in enumerate(records, start=2):
            self._write_cells(sheet, row_index, [
                item['exam_title'],
                item['task_title'],
                item['rank'] if item['rank'] is not None else '',
                item['student_name'],
                item['employee_id'],
                item['department_name'],
                self._format_score(item['score']),
                self._format_score(item['total_score']),
                item['pass_status'],
                item['submission_status'],
                item['time_spent_minutes'] if item['time_spent_minutes'] is not None else '',
            ])

    @staticmethod
    def _format_score(value: Optional[float]) -> str | float:
        if value is None:
            return ''
        if float(value).is_integer():
            return int(value)
        return round(float(value), 1)
